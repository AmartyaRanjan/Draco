"""Group F - file_analysis for files already uploaded through POST /documents (resolved by
document_id; filesystem paths are never accepted from the model or returned to it)."""
import asyncio
import csv
import math
import statistics
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, Field

from app.agent.capabilities.common import check_scope, require, safe_id
from app.agent.contracts import CapabilityError, ErrorCode
from app.agent.registry import Capability

MAX_ROWS = 5000


class AnalyzeIn(BaseModel):
    document_id: str = Field(min_length=1, max_length=128)
    analysis: Literal["overview", "statistics", "sample"] = "overview"
    sheet: str | None = Field(default=None, max_length=31)
    max_sample_rows: int = Field(default=5, ge=1, le=20)


class ColumnInfo(BaseModel):
    name: str
    non_null: int
    numeric: bool
    min: float | None = None
    max: float | None = None
    mean: float | None = None
    stdev: float | None = None


class AnalyzeOut(BaseModel):
    filename: str
    file_type: str
    size_bytes: int
    overview: dict
    columns: list[ColumnInfo] = []
    sample_rows: list[list] = []
    untrusted: bool = True
    notes: list[str] = []
    summary: str = ""


def _num(v):
    try:
        f = float(str(v).replace(",", ""))
        return f if math.isfinite(f) else None
    except (TypeError, ValueError):
        return None


def _table(header, rows, analysis, k):
    cols, notes = [], []
    if len(rows) >= MAX_ROWS:
        notes.append(f"only the first {MAX_ROWS} rows were analyzed")
    if analysis in ("overview", "statistics"):
        for j, name in enumerate(header):
            vals = [r[j] for r in rows if j < len(r) and r[j] not in (None, "")]
            nums = [x for x in (_num(v) for v in vals) if x is not None]
            numeric = bool(vals) and len(nums) >= 0.9 * len(vals)
            cols.append(ColumnInfo(name=str(name), non_null=len(vals), numeric=numeric,
                                   **({"min": min(nums), "max": max(nums), "mean": round(statistics.fmean(nums), 6),
                                       "stdev": round(statistics.stdev(nums), 6) if len(nums) > 1 else 0.0}
                                      if numeric and nums and analysis == "statistics" else {})))
    return cols, [list(r[:len(header)]) for r in rows[:k]], notes


def analyze_path(path: Path, filename: str, analysis: str, sheet: str | None, k: int) -> AnalyzeOut:
    ext = path.suffix.lower().lstrip(".")
    size = path.stat().st_size
    if ext in ("csv", "tsv"):
        with path.open(newline="", encoding="utf-8-sig", errors="replace") as fh:
            rd = csv.reader(fh, delimiter="\t" if ext == "tsv" else ",")
            header = next(rd, [])
            rows = [r for _, r in zip(range(MAX_ROWS), rd)]
        cols, sample, notes = _table(header, rows, analysis, k)
        return AnalyzeOut(filename=filename, file_type=ext, size_bytes=size, overview={"columns": len(header), "rows_read": len(rows)},
                          columns=cols, sample_rows=sample, notes=notes, summary=f"{len(rows)} rows x {len(header)} columns")
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            names = wb.sheetnames
            ws = wb[sheet] if sheet and sheet in names else wb[names[0]]
            it = ws.iter_rows(values_only=True)
            header = [("" if c is None else str(c)) for c in next(it, ())]
            rows = [list(r) for _, r in zip(range(MAX_ROWS), it)]
        finally:
            wb.close()
        cols, sample, notes = _table(header, rows, analysis, k)
        return AnalyzeOut(filename=filename, file_type=ext, size_bytes=size,
                          overview={"sheets": names, "analyzed_sheet": ws.title, "columns": len(header), "rows_read": len(rows)},
                          columns=cols, sample_rows=sample, notes=notes + ["cached formula results were read, formulas not executed"],
                          summary=f"{len(rows)} rows x {len(header)} columns in '{ws.title}'")
    if ext == "pdf":
        from pypdf import PdfReader
        rd = PdfReader(str(path))
        n = len(rd.pages)
        texts = [(rd.pages[i].extract_text() or "") for i in range(min(n, 20))]
        meta = rd.metadata or {}
        return AnalyzeOut(filename=filename, file_type=ext, size_bytes=size,
                          overview={"pages": n, "title": str(meta.get("/Title") or "")[:200], "encrypted": rd.is_encrypted,
                                    "chars_first_pages": [len(t) for t in texts], "first_page_excerpt": (texts[0] if texts else "")[:600]},
                          notes=["pages with ~0 chars likely need OCR"] if any(len(t) < 20 for t in texts) else [],
                          summary=f"{n} pages")
    if ext == "docx":
        from docx import Document
        d = Document(str(path))
        heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")][:30]
        return AnalyzeOut(filename=filename, file_type=ext, size_bytes=size,
                          overview={"paragraphs": len(d.paragraphs), "tables": len(d.tables), "headings": heads,
                                    "first_paragraph": next((p.text for p in d.paragraphs if p.text.strip()), "")[:400]},
                          summary=f"{len(d.paragraphs)} paragraphs, {len(d.tables)} tables")
    if ext in ("txt", "md"):
        t = path.read_text(encoding="utf-8", errors="replace")[:200_000]
        return AnalyzeOut(filename=filename, file_type=ext, size_bytes=size,
                          overview={"chars": len(t), "lines": t.count("\n") + 1, "excerpt": t[:600]}, summary=f"{len(t)} chars")
    raise CapabilityError(ErrorCode.invalid_input, f"File type '.{ext}' is not supported for analysis (csv, tsv, xlsx, pdf, docx, txt, md).")


async def analyze_uploaded_file(a: AnalyzeIn, ctx):
    svc = ctx.services
    lookup = require(svc.document_lookup, "Document registry")
    check_scope(ctx, a.document_id)
    doc = lookup(safe_id(a.document_id), ctx.user_id)
    if doc is None:
        raise CapabilityError(ErrorCode.not_found, "Document not found.")
    root = Path(svc.extra.get("upload_dir", "")).resolve() if svc.extra.get("upload_dir") else None
    path = Path(doc["file_path"]).resolve()
    if root is None or root not in path.parents or not path.is_file():  # only files inside the upload dir
        raise CapabilityError(ErrorCode.not_found, "File is not available for analysis.")
    if path.stat().st_size > svc.settings.analysis_max_bytes:
        raise CapabilityError(ErrorCode.limit_exceeded, "File is too large to analyze.")
    try:
        return await asyncio.to_thread(analyze_path, path, doc["filename"], a.analysis, a.sheet, a.max_sample_rows)
    except CapabilityError:
        raise
    except Exception as exc:
        raise CapabilityError(ErrorCode.invalid_input, f"File could not be parsed ({type(exc).__name__}).") from exc


CAPS = [Capability("analyze_uploaded_file", "file_analysis", "Analyze an uploaded csv/tsv/xlsx/pdf/docx/txt file by document_id: "
                   "structure, column statistics, samples.", AnalyzeIn, AnalyzeOut, analyze_uploaded_file,
                   cost="medium", untrusted_output=True, timeout_s=60)]
